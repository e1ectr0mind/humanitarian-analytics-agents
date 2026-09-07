"""Read project documentation (logframes, proposals) into LLM context.

This is the deliberate exception to the PII boundary (spec §5): files under
project_docs/ are MEANT to be read by the model. Data files never live here.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

SUPPORTED_DOC_SUFFIXES = {".md", ".txt", ".pdf", ".docx", ".xlsx", ".csv"}
_XLSX_ROW_CAP = 200


class DocError(ValueError):
    """Unknown, unsupported, or out-of-tree document request."""


def list_docs(docs_dir: Path) -> list[dict]:
    if not docs_dir.is_dir():
        return []
    return [
        {
            "name": p.name,
            "type": p.suffix.lstrip(".").lower(),
            "size_kb": round(p.stat().st_size / 1024, 1),
        }
        for p in sorted(docs_dir.iterdir())
        if p.is_file() and p.suffix.lower() in SUPPORTED_DOC_SUFFIXES
    ]


def _resolve(docs_dir: Path, name: str) -> Path:
    path = (docs_dir / name).resolve()
    if docs_dir.resolve() not in path.parents:
        raise DocError(f"Document path escapes project_docs/: {name}")
    if not path.is_file():
        raise DocError(f"Document not found: {name}")
    if path.suffix.lower() not in SUPPORTED_DOC_SUFFIXES:
        raise DocError(f"Unsupported document type: {path.suffix}")
    return path


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = io.StringIO()
    for ws in wb.worksheets:
        out.write(f"## sheet: {ws.title}\n")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= _XLSX_ROW_CAP:
                out.write("[... more rows omitted ...]\n")
                break
            out.write("\t".join("" if v is None else str(v) for v in row) + "\n")
    return out.getvalue()


def _read_csv(path: Path) -> str:
    out = io.StringIO()
    with path.open(newline="", encoding="utf-8", errors="replace") as fh:
        for i, row in enumerate(csv.reader(fh)):
            if i >= _XLSX_ROW_CAP:
                out.write("[... more rows omitted ...]\n")
                break
            out.write("\t".join(row) + "\n")
    return out.getvalue()


_READERS = {
    ".md": lambda p: p.read_text(encoding="utf-8", errors="replace"),
    ".txt": lambda p: p.read_text(encoding="utf-8", errors="replace"),
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".xlsx": _read_xlsx,
    ".csv": _read_csv,
}


def read_doc(docs_dir: Path, name: str, max_chars: int = 30000) -> str:
    path = _resolve(docs_dir, name)
    reader = _READERS[path.suffix.lower()]
    try:
        text = reader(path)
    except DocError:
        raise
    except Exception as exc:
        raise DocError(f"Failed to read {name}: {exc}") from exc
    if len(text) > max_chars:
        text = text[:max_chars] + "\n[... document truncated ...]"
    return text
