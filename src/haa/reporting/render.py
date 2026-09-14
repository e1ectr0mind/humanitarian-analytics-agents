"""Render report results to markdown and xlsx (openpyxl)."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from haa.reporting.engine import ReportRun
from haa.reporting.fivew import FiveWTable
from haa.reporting.sources import Period, SourceInfo

SUMMARY_MD_HEADER = ["Code", "Indicator (uk)", "Indicator (en)", "Target", "Actual", "Progress"]
SUMMARY_XLSX_HEADER = [
    "Code", "Indicator (uk)", "Indicator (en)", "Target", "Unit",
    "Actual", "Progress %", "Status", "Reason",
]
SOURCE_HEADER = [
    "Dataset", "File", "Copy", "Rows total", "Rows in scope",
    "Excluded by period", "Rows without a valid date",
]
BREAKDOWN_XLSX_HEADER = ["Code", "Category", "Value", "Note"]

_INVALID_TITLE = re.compile(r"[\[\]:*?/\\]")


def report_paths(reports_dir: Path, kind: str, today: date) -> tuple[Path, Path]:
    stem = f"{kind}_{today.isoformat()}"
    return reports_dir / f"{stem}.md", reports_dir / f"{stem}.xlsx"


def format_value(value: float | None, unit: str | None = None) -> str:
    if value is None:
        return "—"
    text = str(int(value)) if float(value).is_integer() else f"{value:.1f}"
    return f"{text}%" if unit == "percent" else text


def period_label(period: Period | None) -> str:
    return period.label() if period is not None else "all records"


def _cell(value: object) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


def _md_table(header: list[str], rows: list[list[object]]) -> list[str]:
    lines = [
        "| " + " | ".join(_cell(h) for h in header) + " |",
        "|" + "---|" * len(header),
    ]
    lines += ["| " + " | ".join(_cell(v) for v in row) + " |" for row in rows]
    return lines


def _xl(value: object) -> object:
    if isinstance(value, float):
        return int(value) if value.is_integer() else round(value, 4)
    if isinstance(value, str):
        # control characters from the data would make openpyxl refuse the cell
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _sheet_title(name: str, taken: set[str]) -> str:
    clean = ILLEGAL_CHARACTERS_RE.sub("", _INVALID_TITLE.sub("", str(name)))
    base = clean.strip()[:31] or "Sheet"
    title, n = base, 2
    while title.lower() in {t.lower() for t in taken}:
        suffix = f"_{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    taken.add(title)
    return title


def _write_table(ws, header: list[str], rows: list[list[object]]) -> None:
    ws.append([_xl(h) for h in header])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_xl(v) for v in row])


def _source_rows(sources: dict[str, SourceInfo]) -> list[list[object]]:
    return [
        [
            s.dataset, s.path.name, s.copy_label, s.rows_total,
            s.rows_in_scope, s.rows_excluded_by_period, s.rows_bad_date,
        ]
        for s in sources.values()
    ]


def render_indicator_report(
    run: ReportRun, md_path: Path, xlsx_path: Path, generated: date
) -> None:
    lines = [
        "# Indicator progress report",
        "",
        f"Generated: {generated.isoformat()}",
        f"Period: {period_label(run.period)}",
        "",
    ]
    if run.sources:
        lines += ["## Sources", "", *_md_table(SOURCE_HEADER, _source_rows(run.sources)), ""]
        for source in run.sources.values():
            if source.stale_warning:
                lines += [f"> **Warning:** {source.stale_warning}", ""]
            if source.rows_in_scope == 0:
                lines += [
                    f"> **Warning:** no rows in scope for {source.dataset} — its indicators "
                    "are 0 or not computable.",
                    "",
                ]
    summary = [
        [
            r.code, r.name.get("uk", ""), r.name.get("en", ""),
            format_value(r.target, r.unit), format_value(r.actual, r.unit),
            format_value(r.progress_pct, "percent"),
        ]
        for r in run.results
    ]
    lines += ["## Summary", "", *_md_table(SUMMARY_MD_HEADER, summary), ""]
    if run.not_computable:
        lines += ["## Not computable", ""]
        lines += [f"- **{r.code}** — {r.reason}" for r in run.not_computable]
        lines.append("")
    detailed = [r for r in run.computed if r.breakdowns or r.breakdown_errors]
    if detailed:
        lines += ["## Disaggregation", ""]
        for r in detailed:
            lines += [f"### {r.code} — {r.name.get('en') or r.name.get('uk', '')}", ""]
            for dimension, rows in r.breakdowns.items():
                table_rows = [
                    [
                        row.category,
                        format_value(row.value, r.unit)
                        if row.reason is None
                        else f"— ({row.reason})",
                    ]
                    for row in rows
                ]
                lines += [f"**{dimension}**", "", *_md_table([dimension, "Value"], table_rows), ""]
            for dimension, reason in r.breakdown_errors.items():
                lines += [f"**{dimension}**: not available — {reason}", ""]
    md_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text("\n".join(lines), encoding="utf-8")

    wb = Workbook()
    taken: set[str] = set()
    summary_ws = wb.active
    summary_ws.title = _sheet_title("Summary", taken)
    _write_table(
        summary_ws,
        SUMMARY_XLSX_HEADER,
        [
            [
                r.code, r.name.get("uk", ""), r.name.get("en", ""), r.target, r.unit,
                r.actual, r.progress_pct, r.status, r.reason,
            ]
            for r in run.results
        ],
    )
    summary_ws.freeze_panes = "A2"
    about = wb.create_sheet(_sheet_title("About", taken))
    about.append(["Generated", generated.isoformat()])
    about.append(["Period", _xl(period_label(run.period))])
    about.append([None])
    _write_table(about, SOURCE_HEADER, _source_rows(run.sources))
    for source in run.sources.values():
        if source.stale_warning:
            about.append([_xl("Warning"), _xl(source.stale_warning)])
    breakdown_errors = [
        [f"Breakdown not computed: {r.code} by {dimension}", reason]
        for r in run.computed
        for dimension, reason in r.breakdown_errors.items()
    ]
    if breakdown_errors:
        about.append([None])
        for row in breakdown_errors:
            about.append([_xl(v) for v in row])
    by_dimension: dict[str, list[list[object]]] = {}
    for r in run.computed:
        for dimension, rows in r.breakdowns.items():
            by_dimension.setdefault(dimension, []).extend(
                [r.code, row.category, row.value, row.reason] for row in rows
            )
    for dimension, rows in by_dimension.items():
        ws = wb.create_sheet(_sheet_title(dimension, taken))
        _write_table(ws, BREAKDOWN_XLSX_HEADER, rows)
        ws.freeze_panes = "A2"
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def render_5w(
    table: FiveWTable,
    source: SourceInfo,
    period: Period | None,
    md_path: Path,
    xlsx_path: Path,
    generated: date,
    md_row_cap: int = 200,
) -> None:
    copy = source.copy_label
    lines = [
        "# 5W report",
        "",
        f"Generated: {generated.isoformat()}",
        f"Dataset: {source.dataset} ({source.path.name}, {copy} copy)",
        f"Period: {period_label(period)}",
        f"Rows in scope: {source.rows_in_scope} (excluded by period: "
        f"{source.rows_excluded_by_period}; rows without a valid date: {source.rows_bad_date})",
    ]
    if source.stale_warning:
        lines += ["", f"> **Warning:** {source.stale_warning}"]
    lines.append("")
    if not table.rows:
        lines += ["> **Warning:** no activity rows in scope — the table is empty.", ""]
    shown = table.rows[:md_row_cap]
    lines += _md_table(table.columns, [[row[c] for c in table.columns] for row in shown])
    lines.append("")
    if len(table.rows) > md_row_cap:
        lines += [
            f"_Showing the first {md_row_cap} of {len(table.rows)} rows — "
            "the .xlsx has all of them._",
            "",
        ]
    lines += [
        f"**Unique reach (distinct {table.id_field}):** {table.unique_reach}",
        "",
        "A beneficiary with several activities appears in several rows, so the "
        "Beneficiaries column does not add up to the unique reach.",
        "",
        f"**Records without any activity:** {table.rows_without_activity}",
    ]
    md_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text("\n".join(lines), encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "5W"
    _write_table(ws, table.columns, [[row[c] for c in table.columns] for row in table.rows])
    ws.freeze_panes = "A2"
    about = wb.create_sheet("About")
    about_rows = [
        ("Generated", generated.isoformat()),
        ("Period", period_label(period)),
        ("Dataset", source.dataset),
        ("File", source.path.name),
        ("Copy", copy),
        ("Rows in scope", source.rows_in_scope),
        ("Excluded by period", source.rows_excluded_by_period),
        ("Rows without a valid date", source.rows_bad_date),
        (f"Unique reach (distinct {table.id_field})", table.unique_reach),
        ("Records without any activity", table.rows_without_activity),
    ]
    if source.stale_warning:
        about_rows.append(("Warning", source.stale_warning))
    for label, value in about_rows:
        about.append([_xl(label), _xl(value)])
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
