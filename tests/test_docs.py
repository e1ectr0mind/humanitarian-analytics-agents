from pathlib import Path

import openpyxl
import pytest
from docx import Document

from haa.core.tools.docs import DocError, list_docs, read_doc


@pytest.fixture()
def docs_dir(tmp_path: Path) -> Path:
    d = tmp_path / "project_docs"
    d.mkdir()
    (d / "logframe.md").write_text("# Logframe\nIndicator 1.1 target 2500", encoding="utf-8")
    doc = Document()
    doc.add_paragraph("Project proposal narrative.")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Indicator"
    doc.save(d / "proposal.docx")
    wb = openpyxl.Workbook()
    wb.active.append(["indicator", "target"])
    wb.active.append(["1.1", 2500])
    wb.save(d / "targets.xlsx")
    (d / "ignore.exe").write_bytes(b"MZ")
    return d


def test_list_docs(docs_dir: Path) -> None:
    names = {d["name"] for d in list_docs(docs_dir)}
    assert names == {"logframe.md", "proposal.docx", "targets.xlsx"}
    assert all("size_kb" in d and "type" in d for d in list_docs(docs_dir))


def test_read_markdown(docs_dir: Path) -> None:
    assert "target 2500" in read_doc(docs_dir, "logframe.md")


def test_read_docx_paragraphs_and_tables(docs_dir: Path) -> None:
    text = read_doc(docs_dir, "proposal.docx")
    assert "Project proposal narrative." in text
    assert "Indicator" in text


def test_read_xlsx_as_rows(docs_dir: Path) -> None:
    text = read_doc(docs_dir, "targets.xlsx")
    assert "indicator" in text and "2500" in text


def test_max_chars_cap(docs_dir: Path) -> None:
    (docs_dir / "big.txt").write_text("word " * 100_000, encoding="utf-8")
    text = read_doc(docs_dir, "big.txt", max_chars=500)
    assert len(text) <= 550  # cap + truncation notice


def test_unknown_doc_raises(docs_dir: Path) -> None:
    with pytest.raises(DocError, match="not found"):
        read_doc(docs_dir, "missing.md")


def test_unsupported_type_raises(docs_dir: Path) -> None:
    with pytest.raises(DocError, match="[Uu]nsupported"):
        read_doc(docs_dir, "ignore.exe")


def test_path_traversal_rejected(docs_dir: Path, tmp_path: Path) -> None:
    (tmp_path / "outside.md").write_text("secret", encoding="utf-8")
    with pytest.raises(DocError):
        read_doc(docs_dir, "../outside.md")


def test_corrupt_docx_raises_docerror(docs_dir: Path) -> None:
    (docs_dir / "broken.docx").write_bytes(b"this is not a zip archive")
    with pytest.raises(DocError, match="Failed to read"):
        read_doc(docs_dir, "broken.docx")
