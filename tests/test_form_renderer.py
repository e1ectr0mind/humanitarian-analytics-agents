from pathlib import Path

import openpyxl
import pytest

from haa.forms.renderer import render_xlsx
from tests.test_form_model import VALID


@pytest.fixture()
def rendered(tmp_path: Path) -> Path:
    return render_xlsx(VALID, tmp_path / "pdm.xlsx")


def _rows(path: Path, sheet: str) -> list[list]:
    wb = openpyxl.load_workbook(path)
    return [[v if v is not None else "" for v in r] for r in wb[sheet].iter_rows(values_only=True)]


def test_sheets_present(rendered: Path) -> None:
    assert set(openpyxl.load_workbook(rendered).sheetnames) == {"survey", "choices", "settings"}


def test_survey_header_exact(rendered: Path) -> None:
    assert _rows(rendered, "survey")[0] == [
        "type",
        "name",
        "label::Українська (uk)",
        "label::English (en)",
        "hint::Українська (uk)",
        "hint::English (en)",
        "required",
        "relevant",
        "constraint",
    ]


def test_groups_wrap_questions(rendered: Path) -> None:
    types = [r[0] for r in _rows(rendered, "survey")[1:]]
    assert types == [
        "begin_group",
        "select_one yesno",
        "end_group",
        "begin_group",
        "select_one sex",
        "integer",
        "end_group",
    ]


def test_question_row_content(rendered: Path) -> None:
    rows = {r[1]: r for r in _rows(rendered, "survey")[1:]}
    consent = rows["consent_given"]
    assert consent[2] == "Чи згодні ви?" and consent[3] == "Do you consent?"
    assert consent[6] == "yes"  # required
    age = rows["head_age"]
    assert age[4] == "Повних років" and age[5] == "Full years"
    assert age[6] == ""  # not required
    assert age[8] == ". >= 0 and . <= 120"
    assert rows["head_sex"][7] == "${consent_given} = 'yes'"


def test_choices_sheet(rendered: Path) -> None:
    rows = _rows(rendered, "choices")
    assert rows[0] == ["list_name", "name", "label::Українська (uk)", "label::English (en)"]
    assert ["yesno", "yes", "Так", "Yes"] in [list(r) for r in rows[1:]]
    assert ["sex", "female", "Жіноча", "Female"] in [list(r) for r in rows[1:]]


def test_settings_sheet(rendered: Path) -> None:
    header, data = _rows(rendered, "settings")[:2]
    assert header == ["form_title", "form_id", "version", "default_language"]
    assert data[0] == "Survey" and data[1] == "pdm_2026"
    assert len(str(data[2])) == 12 and str(data[2]).isdigit()  # YYYYMMDDHHMM
    assert data[3] == "Українська (uk)"


def test_render_is_overwrite_not_append(tmp_path: Path) -> None:
    path = tmp_path / "pdm.xlsx"
    render_xlsx(VALID, path)
    render_xlsx(VALID, path)
    assert len(_rows(path, "survey")) == 8  # header + 7 rows, not doubled
