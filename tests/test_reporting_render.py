from datetime import date
from pathlib import Path

import openpyxl

from haa.reporting.engine import BreakdownRow, IndicatorResult, ReportRun
from haa.reporting.fivew import FiveWTable
from haa.reporting.measures import MISSING
from haa.reporting.render import (
    format_value,
    render_5w,
    render_indicator_report,
    report_paths,
)
from haa.reporting.sources import Period, SourceInfo

GENERATED = date(2026, 9, 14)
PERIOD = Period("submission_date", date(2026, 6, 1), date(2026, 8, 31))
SOURCE = SourceInfo(
    dataset="beneficiaries",
    path=Path("data/beneficiaries_clean.xlsx"),
    used_clean=True,
    rows_total=3030,
    rows_in_scope=3000,
    rows_excluded_by_period=30,
    rows_bad_date=2,
)
RUN = ReportRun(
    period=PERIOD,
    sources={"beneficiaries": SOURCE},
    results=[
        IndicatorResult(
            code="1.1",
            name={"uk": "Охоплені", "en": "Reached"},
            dataset="beneficiaries",
            target=2500.0,
            unit="households",
            actual=3000.0,
            progress_pct=120.0,
            status="computed",
            breakdowns={
                "oblast": [BreakdownRow("Донецька", 400.0), BreakdownRow(MISSING, 5.0)],
            },
            breakdown_errors={
                "sex": "column 'sex' not found in the dataset — check profile_dataset",
            },
        ),
        IndicatorResult(
            code="1.2",
            name={"uk": "Частка", "en": "Share | female"},
            dataset="beneficiaries",
            target=55.0,
            unit="percent",
            actual=57.63,
            progress_pct=104.8,
            status="computed",
        ),
        IndicatorResult(
            code="2.1",
            name={"uk": "Грошова", "en": "Cash"},
            dataset=None,
            target=1200.0,
            unit="households",
            actual=None,
            progress_pct=None,
            status="not_computable",
            reason="no measure block — add one to the registry (designer)",
        ),
    ],
)
TABLE = FiveWTable(
    columns=["Organization", "oblast", "Period", "Activity", "Beneficiaries"],
    rows=[
        {"Organization": "IMC", "oblast": "X", "Period": "2026-06", "Activity": "cash",
         "Beneficiaries": 1},
        {"Organization": "IMC", "oblast": "X", "Period": "2026-06", "Activity": "health",
         "Beneficiaries": 2},
        {"Organization": "IMC", "oblast": "Y", "Period": "2026-07", "Activity": "cash",
         "Beneficiaries": 1},
    ],
    id_field="_uuid",
    unique_reach=2,
    rows_without_activity=1,
)


def _paths(tmp_path: Path, kind: str) -> tuple[Path, Path]:
    return report_paths(tmp_path / "reports", kind, GENERATED)


def _rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_report_paths(tmp_path: Path) -> None:
    md, xlsx = report_paths(tmp_path, "indicators", GENERATED)
    assert md.name == "indicators_2026-09-14.md" and xlsx.name == "indicators_2026-09-14.xlsx"


def test_format_value() -> None:
    assert format_value(3000.0, "households") == "3000"
    assert format_value(57.63, "percent") == "57.6%"
    assert format_value(120.0, "percent") == "120%"
    assert format_value(None) == "—"


def test_indicator_markdown(tmp_path: Path) -> None:
    md, xlsx = _paths(tmp_path, "indicators")
    render_indicator_report(RUN, md, xlsx, GENERATED)
    text = md.read_text(encoding="utf-8")
    assert "Period: 2026-06-01 .. 2026-08-31 (by submission_date)" in text
    assert "| beneficiaries | beneficiaries_clean.xlsx | clean | 3030 | 3000 | 30 | 2 |" in text
    assert "| 1.1 | Охоплені | Reached | 2500 | 3000 | 120% |" in text
    assert "| 1.2 | Частка | Share \\| female | 55% | 57.6% | 104.8% |" in text
    assert "| 2.1 | Грошова | Cash | 1200 | — | — |" in text
    assert "- **2.1** — no measure block" in text
    assert f"| {MISSING} | 5 |" in text
    assert "**sex**: not available — column 'sex' not found" in text


def test_indicator_markdown_without_period_and_empty_scope(tmp_path: Path) -> None:
    empty = SourceInfo("beneficiaries", Path("b.xlsx"), False, 10, 0, 10, 0)
    run = ReportRun(period=None, sources={"beneficiaries": empty}, results=RUN.results)
    md, xlsx = _paths(tmp_path, "indicators")
    render_indicator_report(run, md, xlsx, GENERATED)
    text = md.read_text(encoding="utf-8")
    assert "Period: all records" in text
    assert "no rows in scope for beneficiaries" in text


def test_indicator_xlsx(tmp_path: Path) -> None:
    md, xlsx = _paths(tmp_path, "indicators")
    render_indicator_report(RUN, md, xlsx, GENERATED)
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["Summary", "About", "oblast"]
    summary = _rows(wb["Summary"])
    assert summary[0] == [
        "Code", "Indicator (uk)", "Indicator (en)", "Target", "Unit",
        "Actual", "Progress %", "Status", "Reason",
    ]
    assert summary[1] == ["1.1", "Охоплені", "Reached", 2500, "households", 3000, 120,
                          "computed", None]
    assert summary[2][5] == 57.63
    assert summary[3][5] is None and summary[3][7] == "not_computable"
    about = _rows(wb["About"])
    assert about[0][:2] == ["Generated", "2026-09-14"]
    assert about[1][1] == "2026-06-01 .. 2026-08-31 (by submission_date)"
    assert about[4] == ["beneficiaries", "beneficiaries_clean.xlsx", "clean", 3030, 3000, 30, 2]
    assert wb["About"]["A4"].font.bold
    assert _rows(wb["oblast"])[1:] == [["1.1", "Донецька", 400, None], ["1.1", MISSING, 5, None]]


def test_sheet_titles_are_sanitized_and_unique(tmp_path: Path) -> None:
    result = RUN.results[0]
    run = ReportRun(
        period=None,
        sources={},
        results=[
            IndicatorResult(
                code="1", name=result.name, dataset="b", target=None, unit=None, actual=1.0,
                progress_pct=None, status="computed",
                breakdowns={"a/b:c": [BreakdownRow("x", 1.0)], "summary": [BreakdownRow("y", 1.0)]},
            )
        ],
    )
    md, xlsx = _paths(tmp_path, "indicators")
    render_indicator_report(run, md, xlsx, GENERATED)
    assert openpyxl.load_workbook(xlsx).sheetnames == ["Summary", "About", "abc", "summary_2"]


def test_5w_markdown(tmp_path: Path) -> None:
    md, xlsx = _paths(tmp_path, "5w")
    render_5w(TABLE, SOURCE, PERIOD, md, xlsx, GENERATED)
    text = md.read_text(encoding="utf-8")
    assert "Dataset: beneficiaries (beneficiaries_clean.xlsx, clean copy)" in text
    assert "| Organization | oblast | Period | Activity | Beneficiaries |" in text
    assert "| IMC | X | 2026-06 | health | 2 |" in text
    assert "**Unique reach (distinct _uuid):** 2" in text
    assert "**Records without any activity:** 1" in text


def test_5w_markdown_row_cap(tmp_path: Path) -> None:
    md, xlsx = _paths(tmp_path, "5w")
    render_5w(TABLE, SOURCE, None, md, xlsx, GENERATED, md_row_cap=2)
    text = md.read_text(encoding="utf-8")
    assert "| IMC | Y | 2026-07 | cash | 1 |" not in text
    assert "first 2 of 3 rows" in text
    assert len(_rows(openpyxl.load_workbook(xlsx)["5W"])) == 4  # the xlsx keeps every row


def test_5w_xlsx(tmp_path: Path) -> None:
    md, xlsx = _paths(tmp_path, "5w")
    render_5w(TABLE, SOURCE, PERIOD, md, xlsx, GENERATED)
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["5W", "About"]
    rows = _rows(wb["5W"])
    assert rows[0] == TABLE.columns
    assert rows[2] == ["IMC", "X", "2026-06", "health", 2]
    about = {r[0]: r[1] for r in _rows(wb["About"])}
    assert about["Unique reach (distinct _uuid)"] == 2
    assert about["Copy"] == "clean"


def test_5w_empty_table_warns(tmp_path: Path) -> None:
    empty = FiveWTable(columns=TABLE.columns, rows=[], id_field="_uuid", unique_reach=0,
                       rows_without_activity=0)
    md, xlsx = _paths(tmp_path, "5w")
    render_5w(empty, SOURCE, None, md, xlsx, GENERATED)
    assert "table is empty" in md.read_text(encoding="utf-8")
    assert _rows(openpyxl.load_workbook(xlsx)["5W"]) == [TABLE.columns]
